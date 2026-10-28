from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PptInches, Pt as PptPt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

NAVY = "0B132B"
INK = "182230"
TEAL = "0B9F9A"
CYAN = "28D7C3"
ORANGE = "E78E42"
LIGHT = "F3F6FA"
MID = "65738A"
LINE = "D9E0E8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, name="Microsoft YaHei", size=11, color=INK, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def style_paragraph(paragraph, space_after=6, line=1.15):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line


def add_doc_paragraph(doc, text="", style=None, size=11, color=INK, bold=False, after=6, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold)
    style_paragraph(p, after)
    if align is not None:
        p.alignment = align
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    style_paragraph(p, 3, 1.12)
    return p


def add_doc_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, head in enumerate(headers):
        hdr[idx].text = ""
        set_cell_shading(hdr[idx], NAVY)
        set_cell_border(hdr[idx])
        hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(head)
        set_run_font(r, size=10, color=WHITE, bold=True)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cells[cidx].text = ""
            set_cell_border(cells[cidx])
            set_cell_shading(cells[cidx], "F7F9FC" if ridx % 2 else WHITE)
            cells[cidx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[cidx].paragraphs[0]
            r = p.add_run(str(value))
            set_run_font(r, size=9.5, color=INK)
            style_paragraph(p, 2, 1.08)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in (("Title", 25, "000000"), ("Heading 1", 17, "000000"), ("Heading 2", 13, "000000")):
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True

    # Header and footer
    header = sec.header.paragraphs[0]
    header.text = "SecureRelayChat  ·  密码学知识总结"
    for r in header.runs:
        set_run_font(r, size=9, color=MID)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("教育用途演示  |  基于项目源码整理")
    set_run_font(r, size=8.5, color=MID)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("SecureRelayChat 密码学知识总结")
    set_run_font(r, size=25, color="000000", bold=True)
    style_paragraph(p, 4, 1.0)
    p = add_doc_paragraph(doc, "从密钥协商到消息认证的教育用途双人端到端加密演示", size=13, color="000000", bold=True, after=14)
    add_doc_paragraph(doc, "这份材料对应项目中的实际实现，目标是帮助读者理解四个核心问题：双方怎样得到共同密钥，消息怎样保密，怎样确认消息确实来自对方，以及怎样降低中间人替换公钥的风险。", size=11.5, color=INK, after=10)
    add_doc_table(doc, ["密码学层", "项目使用", "解决的问题"], [
        ["密钥协商", "X25519 + HKDF-SHA256", "双方在不发送会话密钥的情况下得到同一把密钥"],
        ["消息加密", "AES-256-GCM", "隐藏明文，并检测密文是否被篡改"],
        ["身份认证", "Ed25519", "证明消息由持有对应身份私钥的一方签发"],
        ["人工防中间人", "6 位安全码", "让用户比较双方身份材料是否一致"],
    ], widths=[1.2, 2.0, 3.7])

    doc.add_heading("一 项目要解决的安全问题", level=1)
    add_doc_paragraph(doc, "项目把服务器视为不可信中继。服务器知道谁把数据发给谁，也可以记录或转发网络帧，但不应该知道消息明文，不能替客户端完成解密。真正的密码学操作全部在两个客户端本地完成。", size=10.8)
    add_bullet(doc, "保密性：服务器看到的是随机 nonce 和密文，不能直接读出消息。")
    add_bullet(doc, "完整性：密文或附加认证数据被改动时，AES-GCM 解密失败。")
    add_bullet(doc, "来源认证：Ed25519 签名让接收方确认消息由持有对方身份私钥的一方产生。")
    add_bullet(doc, "防中间人：用户通过可信渠道比较安全码，发现公钥被替换。")
    add_doc_paragraph(doc, "这里的关键思想是分层：加密负责‘别人看不懂’，认证负责‘别人不能冒充’，安全码负责‘用户确认眼前的身份确实是那个人’。", size=10.8, color=TEAL, bold=True, after=10)

    doc.add_heading("二 整体密码学链路", level=1)
    add_doc_paragraph(doc, "一次会话可以概括为下面这条链路。每一层只承担一个清晰职责，组合起来形成端到端保护。", size=10.8)
    add_doc_table(doc, ["阶段", "发生了什么", "输出"], [
        ["身份准备", "客户端生成长期 Ed25519 签名密钥和 X25519 身份交换密钥，并只保存在本地", "身份公钥"],
        ["临时握手", "双方各生成一次性 X25519 临时密钥，并交换临时公钥", "共享秘密 S"],
        ["会话派生", "HKDF 使用共享秘密和双方临时公钥生成 32 字节会话密钥", "session_key"],
        ["安全码确认", "身份公钥经过固定顺序归一化后计算 6 位数字", "safety number"],
        ["消息发送", "AES-GCM 加密，Ed25519 对 nonce、密文和 AAD 的摘要签名", "消息 envelope"],
        ["消息接收", "先验签，再用 GCM 解密，失败则拒绝消息", "明文或错误"],
    ], widths=[1.1, 4.5, 1.3])

    doc.add_heading("三 X25519 与 HKDF 密钥协商", level=1)
    doc.add_heading("X25519 的思想", level=2)
    add_doc_paragraph(doc, "X25519 是椭圆曲线 Diffie-Hellman 的现代实现。Alice 和 Bob 各自保留一个私钥，只发送对应的公钥。两边分别用‘自己的私钥 + 对方的公钥’计算，得到相同的共享秘密；旁观者只看到公钥，无法直接算出这个秘密。", size=10.8)
    add_doc_paragraph(doc, "项目中使用的是临时 X25519 密钥。每次聊天会话重新生成临时私钥，私钥只在内存中存在，退出时显式清空。", size=10.8)
    doc.add_heading("HKDF 的思想", level=2)
    add_doc_paragraph(doc, "X25519 得到的共享秘密更像一段原始材料，不直接拿来做 AES 密钥。HKDF-SHA256 负责把它‘提炼’成固定长度、用途明确的 32 字节会话密钥。项目把双方临时公钥按字典序拼接为 salt，并使用固定 info 文本，确保两边输入一致，也让不同会话的派生结果绑定到具体临时密钥。", size=10.8)
    add_doc_table(doc, ["对象", "项目中的做法", "重要原因"], [
        ["私钥", "只在本地内存使用", "避免把长期秘密交给服务器"],
        ["临时公钥", "随握手帧发送", "让双方得到同一共享秘密"],
        ["salt", "双方临时公钥排序后拼接", "绑定会话，避免不同会话意外复用密钥"],
        ["输出", "32 字节", "匹配 AES-256 的密钥长度"],
    ], widths=[1.1, 3.1, 2.7])

    doc.add_heading("四 AES-256-GCM 消息加密", level=1)
    add_doc_paragraph(doc, "AES 是对称加密算法：加密和解密使用同一把会话密钥。项目选择 AES-256-GCM，因为它把保密性和完整性放在一个接口里处理。发送方输入会话密钥、随机 12 字节 nonce、明文和 AAD，输出密文以及 GCM authentication tag。", size=10.8)
    add_doc_paragraph(doc, "AAD 是‘附加认证数据’，项目把 sender:recipient 放入 AAD。它不被加密，但会参与认证，因此攻击者不能把同一条密文悄悄改成发给另一个用户名。", size=10.8)
    add_bullet(doc, "每条消息重新生成 nonce；同一会话密钥下 nonce 不能重复。")
    add_bullet(doc, "密文被改动、nonce 被改动或 AAD 不匹配时，GCM 校验失败。")
    add_bullet(doc, "接收方拿到的不是‘先解密再检查’，而是由 GCM 同时完成解密与认证。")

    doc.add_heading("五 Ed25519 数字签名与安全码", level=1)
    doc.add_heading("Ed25519 的思想", level=2)
    add_doc_paragraph(doc, "Ed25519 是公钥数字签名算法。发送方用身份私钥对消息摘要签名，接收方用已知的身份公钥验签。签名不负责隐藏消息，而是回答‘这条数据是不是对应身份私钥的持有者签发的，以及内容有没有被改过’。", size=10.8)
    add_doc_paragraph(doc, "项目先计算 SHA-256(nonce + ciphertext + aad)，再对摘要签名。接收时严格先验签，验签失败就停止，不对未认证密文做解密。", size=10.8)
    doc.add_heading("6 位安全码的思想", level=2)
    add_doc_paragraph(doc, "安全码不是新的加密算法，而是把双方长期身份公钥压缩成一个人可以比较的短指纹。项目对双方签名公钥和交换公钥做固定顺序归一化，再取 SHA-256 的前 4 字节映射成 6 位数字，因此 Alice 和 Bob 会得到同一个安全码。", size=10.8)
    add_doc_paragraph(doc, "当服务器替换了 Bob 的公钥时，Alice 看到的安全码会与 Bob 本人看到的不一致。只要用户通过电话、面对面等可信渠道比较安全码，就有机会发现这种中间人攻击。", size=10.8, color=TEAL, bold=True)

    doc.add_heading("六 一条消息如何走完全部链路", level=1)
    add_doc_table(doc, ["顺序", "发送方", "接收方"], [
        ["1", "用 sender:recipient 形成 AAD", "从 envelope 恢复同样的 AAD"],
        ["2", "AES-GCM 生成 nonce 和 ciphertext", "读取 nonce 和 ciphertext"],
        ["3", "对 nonce + ciphertext + AAD 的 SHA-256 摘要签名", "用对方 Ed25519 公钥验签"],
        ["4", "通过 relay 原样发送 envelope", "验签失败则拒绝，不解密"],
        ["5", "服务器只处理 to 和 payload 外层字段", "验签通过后再调用 AES-GCM 解密"],
    ], widths=[0.7, 3.0, 3.2])
    add_doc_paragraph(doc, "这个顺序体现了一个重要安全原则：先确认数据来源和完整性，再把密文交给解密逻辑。", size=10.8, color=ORANGE, bold=True)

    doc.add_heading("七 密钥生命周期与项目边界", level=1)
    add_doc_table(doc, ["密钥材料", "生命周期", "保存位置", "项目处理"], [
        ["Ed25519 身份私钥", "长期", ".keys/<用户名>.json", "不发送到服务器"],
        ["X25519 身份私钥", "长期", ".keys/<用户名>.json", "不发送到服务器"],
        ["X25519 临时私钥", "单次会话", "客户端内存", "/quit 时清空"],
        ["AES 会话密钥", "单次会话", "客户端内存", "/quit 时清空"],
    ], widths=[1.7, 1.4, 2.3, 1.5])
    add_doc_paragraph(doc, "本项目是教学演示，不等同于生产级即时通信协议。它明确不包含 Double Ratchet 前向保密、离线消息、群聊、多设备、历史持久化和完整安全审计。", size=10.8, color=ORANGE, bold=True)

    doc.add_heading("八 汇报时可以这样总结", level=1)
    add_doc_paragraph(doc, "这个项目没有把‘加密’当成单一函数，而是把安全目标拆开：X25519 解决双方如何得到共同秘密，HKDF 把秘密整理成可用的会话密钥，AES-GCM 保护消息内容和密文完整性，Ed25519 证明消息来源，安全码让用户参与身份确认。服务器只负责转发，因此即使服务器不可信，消息明文仍留在两个客户端之间。", size=11.2, color=NAVY, bold=True, after=12)
    add_doc_paragraph(doc, "阅读项目时，建议依次查看 crypto.py、key_store.py、envelope.py、server.py 和 client.py。这样可以从密码学原语一路追踪到握手和 CLI 行为。", size=10.8)

    path = OUT / "SecureRelayChat_密码学知识总结.docx"
    doc.save(path)
    return path


def ppt_color(hexstr):
    return PptRGB(*bytes.fromhex(hexstr))


def add_text(slide, text, x, y, w, h, size=20, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font="Microsoft YaHei"):
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = PptInches(0.04)
    tf.margin_right = PptInches(0.04)
    tf.margin_top = PptInches(0.02)
    tf.margin_bottom = PptInches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = 0
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = PptPt(size)
    r.font.bold = bold
    r.font.color.rgb = ppt_color(color)
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.55, 0.38, 12.0, 0.48, size=27, color=WHITE, bold=True)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.55), PptInches(0.96), PptInches(1.15), PptInches(0.05)).fill.solid()
    line = slide.shapes[-1]
    line.fill.fore_color.rgb = ppt_color(CYAN)
    line.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.08, 11.6, 0.35, size=12.5, color="AAB7CF")


def base_slide(prs, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = ppt_color(NAVY)
    add_text(slide, f"SECURE RELAY CHAT  /  {num:02d}", 0.58, 7.12, 4, 0.2, size=8.5, color="73819A")
    return slide


def rounded_box(slide, text, x, y, w, h, fill="14213D", line="2C4266", color=WHITE, size=16, bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = ppt_color(fill)
    shape.line.color.rgb = ppt_color(line); shape.line.width = PptPt(1)
    shape.text_frame.clear(); shape.text_frame.word_wrap = True
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = shape.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.name = "Microsoft YaHei"; r.font.size = PptPt(size); r.font.bold = bold; r.font.color.rgb = ppt_color(color)
    return shape


def add_bullets_ppt(slide, items, x, y, w, h, size=19, color="DCE6F5", bullet_color=CYAN, spacing=8):
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = PptInches(0.12); tf.margin_right = PptInches(0.05)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.level = 0; p.space_after = PptPt(spacing)
        p.font.name = "Microsoft YaHei"; p.font.size = PptPt(size); p.font.color.rgb = ppt_color(color)
        p._p.get_or_add_pPr().set('marL', '360')
        p._p.get_or_add_pPr().set('indent', '-180')
        p.text = "•  " + item
    return box


def build_pptx():
    prs = Presentation()
    # Use the exact 16:9 canvas expected by the presentation validator.
    prs.slide_width = 12192000
    prs.slide_height = 6858000

    # 1
    s = base_slide(prs, 1)
    add_text(s, "SecureRelayChat", 0.72, 1.05, 9.2, 0.78, size=42, color=WHITE, bold=True)
    add_text(s, "从密钥协商到消息认证的端到端加密演示", 0.76, 2.0, 9.0, 0.55, size=22, color=CYAN, bold=True)
    add_text(s, "密码学知识汇报", 0.78, 2.8, 5.0, 0.35, size=16, color="AAB7CF")
    for x1, y1, x2, y2 in [(8.9,3.1,9.5,2.33),(11.1,3.1,10.6,2.33)]:
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PptInches(x1), PptInches(y1), PptInches(x2), PptInches(y2)); c.line.color.rgb=ppt_color(CYAN); c.line.width=PptPt(1.8)
    rounded_box(s, "不可信中继服务器", 8.7, 1.6, 2.6, 0.72, fill="1A2B49", line="28D7C3", color=CYAN, size=18, bold=True)
    rounded_box(s, "客户端 A", 8.0, 3.1, 1.8, 0.72, fill="183C4A", line="28D7C3", color=WHITE, size=18, bold=True)
    rounded_box(s, "客户端 B", 10.2, 3.1, 1.8, 0.72, fill="183C4A", line="28D7C3", color=WHITE, size=18, bold=True)

    # 2
    s = base_slide(prs, 2); add_title(s, "项目面对的安全问题", "把网络风险拆成四个密码学问题")
    add_bullets_ppt(s, ["服务器可以转发和记录网络帧，但不应该看到消息明文", "攻击者可能篡改密文，也可能替换对方公钥", "客户端需要知道消息来自谁，并确认对方身份没有被冒充"], 0.8, 1.75, 7.0, 3.7, size=20)
    rounded_box(s, "保密性\n别人看不懂", 8.25, 1.65, 3.5, 0.88, fill="183C4A", line=CYAN, size=18, bold=True)
    rounded_box(s, "完整性\n改了就能发现", 8.25, 2.75, 3.5, 0.88, fill="1A2B49", line="6B8DB7", size=18, bold=True)
    rounded_box(s, "来源认证\n知道是谁发的", 8.25, 3.85, 3.5, 0.88, fill="3A2F25", line=ORANGE, size=18, bold=True)
    rounded_box(s, "身份确认\n发现中间人", 8.25, 4.95, 3.5, 0.88, fill="28344D", line="B6C7E5", size=18, bold=True)

    # 3
    s = base_slide(prs, 3); add_title(s, "四层密码学链路", "每个算法只承担一个清晰职责")
    labels = [("X25519", "共同秘密", "183C4A"), ("HKDF-SHA256", "会话密钥", "1A2B49"), ("AES-256-GCM", "加密与完整性", "183C4A"), ("Ed25519", "消息来源", "3A2F25"), ("安全码", "人工确认身份", "28344D")]
    xs = [0.65, 3.05, 5.45, 7.85, 10.25]
    for i,(name,sub,fill) in enumerate(labels):
        rounded_box(s, name + "\n" + sub, xs[i], 2.5, 1.95, 1.25, fill=fill, line=CYAN if i<3 else ORANGE, size=16, bold=True)
        if i < len(labels)-1:
            c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PptInches(xs[i]+1.95), PptInches(3.12), PptInches(xs[i+1]), PptInches(3.12)); c.line.color.rgb=ppt_color("7890B0"); c.line.width=PptPt(1.5)
    add_text(s, "握手阶段", 0.95, 4.15, 4.5, 0.35, size=14, color=CYAN, bold=True)
    add_text(s, "消息阶段", 5.8, 4.15, 3.5, 0.35, size=14, color=ORANGE, bold=True)
    add_text(s, "用户确认阶段", 10.25, 4.15, 2.4, 0.35, size=14, color="B6C7E5", bold=True)
    add_text(s, "服务器只转发 relay 帧，密码学处理全部在客户端完成。", 2.2, 5.35, 8.8, 0.35, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    # 4
    s = base_slide(prs, 4); add_title(s, "身份密钥与临时密钥", "长期身份负责‘是谁’，临时密钥负责‘这一次会话’")
    rounded_box(s, "长期身份\nEd25519 私钥\nX25519 身份私钥", 0.85, 1.8, 3.15, 2.0, fill="1A2B49", line="6B8DB7", size=18, bold=True)
    rounded_box(s, "单次会话\nX25519 临时私钥\n内存中存在", 4.95, 1.8, 3.15, 2.0, fill="183C4A", line=CYAN, size=18, bold=True)
    rounded_box(s, "消息保护\nAES 会话密钥\n退出时清空", 9.05, 1.8, 3.15, 2.0, fill="3A2F25", line=ORANGE, size=18, bold=True)
    for x1,x2 in [(4.0,4.95),(8.1,9.05)]:
        c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PptInches(x1),PptInches(2.8),PptInches(x2),PptInches(2.8)); c.line.color.rgb=ppt_color("7890B0"); c.line.width=PptPt(1.5)
    add_bullets_ppt(s, ["长期私钥保存在本地 .keys 目录，不发送给服务器", "临时私钥每次会话重新生成，避免所有会话共用同一把密钥", "项目没有实现 Double Ratchet，因此不提供生产级前向保密"], 1.0, 4.45, 11.0, 1.6, size=18)

    # 5
    s = base_slide(prs, 5); add_title(s, "X25519 与 HKDF", "双方不直接传输会话密钥，而是各自计算出同一份秘密")
    rounded_box(s, "Alice\n私钥 a\n公钥 A", 0.9, 2.0, 2.15, 1.35, fill="183C4A", line=CYAN, size=19, bold=True)
    rounded_box(s, "Bob\n私钥 b\n公钥 B", 10.3, 2.0, 2.15, 1.35, fill="183C4A", line=CYAN, size=19, bold=True)
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PptInches(3.05),PptInches(2.68),PptInches(10.3),PptInches(2.68)); c.line.color.rgb=ppt_color("7890B0"); c.line.width=PptPt(2)
    add_text(s, "交换 A 与 B", 5.2, 2.18, 2.8, 0.3, size=16, color="B6C7E5", bold=True, align=PP_ALIGN.CENTER)
    rounded_box(s, "a × B = S", 3.15, 3.75, 2.65, 0.82, fill="1A2B49", line="6B8DB7", size=20, bold=True)
    rounded_box(s, "b × A = S", 7.55, 3.75, 2.65, 0.82, fill="1A2B49", line="6B8DB7", size=20, bold=True)
    rounded_box(s, "HKDF-SHA256\nS + salt + info\n32 字节 session_key", 4.55, 5.2, 4.25, 1.05, fill="3A2F25", line=ORANGE, size=18, bold=True)
    add_text(s, "旁观者看到 A、B，但没有 a、b，难以得到 S。", 2.25, 6.42, 8.8, 0.32, size=16, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 6
    s = base_slide(prs, 6); add_title(s, "AES-256-GCM", "一次操作同时提供保密性与密文完整性")
    rounded_box(s, "明文", 0.9, 2.25, 1.75, 0.75, fill="183C4A", line=CYAN, size=20, bold=True)
    rounded_box(s, "AES-GCM\n32 字节密钥\n随机 12 字节 nonce", 3.35, 1.75, 3.1, 1.75, fill="1A2B49", line="6B8DB7", size=18, bold=True)
    rounded_box(s, "密文 + tag", 7.2, 2.25, 2.0, 0.75, fill="3A2F25", line=ORANGE, size=20, bold=True)
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PptInches(2.65),PptInches(2.62),PptInches(3.35),PptInches(2.62)); c.line.color.rgb=ppt_color(CYAN); c.line.width=PptPt(1.8)
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PptInches(6.45),PptInches(2.62),PptInches(7.2),PptInches(2.62)); c.line.color.rgb=ppt_color(ORANGE); c.line.width=PptPt(1.8)
    add_bullets_ppt(s, ["nonce 不是秘密，但同一会话密钥下不能重复", "tag 用来检测密文、nonce 或 AAD 是否被篡改", "项目把 sender:recipient 作为 AAD，防止密文被换收件人"], 1.05, 4.35, 10.8, 1.65, size=18)

    # 7
    s = base_slide(prs, 7); add_title(s, "Ed25519 与六位安全码", "数字签名回答‘谁发的’，安全码帮助用户回答‘是不是那个人’")
    rounded_box(s, "Ed25519 私钥\n签名摘要", 0.9, 1.85, 2.75, 1.2, fill="3A2F25", line=ORANGE, size=18, bold=True)
    rounded_box(s, "消息摘要\nSHA-256(nonce + ciphertext + AAD)", 4.35, 1.85, 4.35, 1.2, fill="1A2B49", line="6B8DB7", size=18, bold=True)
    rounded_box(s, "签名\n64 字节", 9.35, 1.85, 2.25, 1.2, fill="183C4A", line=CYAN, size=18, bold=True)
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PptInches(3.65),PptInches(2.45),PptInches(4.35),PptInches(2.45)); c.line.color.rgb=ppt_color(ORANGE); c.line.width=PptPt(1.8)
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PptInches(8.7),PptInches(2.45),PptInches(9.35),PptInches(2.45)); c.line.color.rgb=ppt_color(CYAN); c.line.width=PptPt(1.8)
    rounded_box(s, "双方长期身份公钥\n归一化后计算 SHA-256\n取 6 位数字", 3.8, 4.05, 5.8, 1.35, fill="183C4A", line=CYAN, size=19, bold=True)
    add_text(s, "通过电话或面对面比较安全码，可发现服务器替换公钥的中间人攻击。", 1.2, 5.95, 10.9, 0.4, size=18, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 8
    s = base_slide(prs, 8); add_title(s, "消息接收顺序", "项目明确要求先验签，再解密")
    steps=[("1", "解析 envelope", "字段和 Base64 长度先校验"), ("2", "验 Ed25519 签名", "失败就拒绝消息"), ("3", "AES-GCM 解密", "tag 失败就拒绝消息"), ("4", "打印明文", "只有通过前面检查才执行")]
    for i,(n,t,d) in enumerate(steps):
        y=1.72+i*1.25
        rounded_box(s, n, 1.0, y, 0.62, 0.62, fill=CYAN if i<2 else ORANGE, line=WHITE, color=NAVY, size=20, bold=True)
        add_text(s, t, 2.0, y+0.02, 3.4, 0.3, size=21, color=WHITE, bold=True)
        add_text(s, d, 6.05, y+0.04, 5.6, 0.28, size=17, color="B6C7E5")
        if i < 3:
            c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,PptInches(1.31),PptInches(y+0.62),PptInches(1.31),PptInches(y+1.18)); c.line.color.rgb=ppt_color("7890B0"); c.line.width=PptPt(1.4)
    add_text(s, "先确认来源和完整性，再把密文交给解密逻辑。", 2.15, 6.45, 8.8, 0.32, size=18, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

    # 9
    s = base_slide(prs, 9); add_title(s, "能保护什么，不能保护什么", "把演示目标和生产级系统区分开")
    rounded_box(s, "本项目做到\n消息保密\n篡改可检测\n消息来源可验证\n用户可发现公钥替换", 0.95, 1.7, 5.2, 3.2, fill="183C4A", line=CYAN, size=19, bold=True)
    rounded_box(s, "本项目没有做到\nDouble Ratchet 前向保密\n离线消息与多设备\n群聊与文件传输\n生产级安全审计", 7.15, 1.7, 5.2, 3.2, fill="3A2F25", line=ORANGE, size=19, bold=True)
    add_text(s, "安全性来自密码学、密钥管理和用户确认共同作用，不能只看某一个算法名称。", 1.1, 5.75, 11.0, 0.4, size=18, color="B6C7E5", bold=True, align=PP_ALIGN.CENTER)

    # 10
    s = base_slide(prs, 10); add_title(s, "汇报总结", "一条主线记住这个项目")
    add_text(s, "协商共同秘密", 0.8, 1.7, 2.1, 0.35, size=19, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "保护消息内容", 3.35, 1.7, 2.1, 0.35, size=19, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "证明消息来源", 5.9, 1.7, 2.1, 0.35, size=19, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "确认真实身份", 8.45, 1.7, 2.1, 0.35, size=19, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    names=["X25519 + HKDF", "AES-256-GCM", "Ed25519", "安全码"]
    xs=[0.8,3.35,5.9,8.45]
    for x,name in zip(xs,names): rounded_box(s,name,x,2.35,2.1,0.85,fill="1A2B49",line=CYAN if x<5.9 else ORANGE,size=17,bold=True)
    add_text(s, "服务器只转发，客户端完成协商、加密、认证与解密。", 1.0, 4.3, 11.2, 0.5, size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "这就是端到端加密演示的核心思想：把信任边界放回通信双方。", 1.15, 5.15, 11.0, 0.4, size=19, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    path = OUT / "SecureRelayChat_密码学链路汇报.pptx"
    prs.save(path)
    return path


if __name__ == "__main__":
    print(build_docx())
    print(build_pptx())
